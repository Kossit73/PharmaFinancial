"""Document-ingestion and lightweight RAG helpers for the Streamlit app."""
from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import List, Mapping

try:  # pragma: no cover - optional dependency for Excel ingestion
    from openpyxl import load_workbook
except Exception:  # pragma: no cover - import guard when package missing
    load_workbook = None  # type: ignore

try:  # pragma: no cover - optional dependency for Word ingestion
    from docx import Document
except Exception:  # pragma: no cover - import guard when package missing
    Document = None  # type: ignore

try:  # pragma: no cover - optional dependency for PDF ingestion
    from PyPDF2 import PdfReader
except Exception:  # pragma: no cover - import guard when package missing
    PdfReader = None  # type: ignore


def load_payload_from_bytes(data: bytes, suffix: str) -> Mapping[str, object]:
    suffix = suffix or ".json"
    if suffix in {".json", ""}:
        return load_payload_from_text(data.decode("utf-8"))
    if suffix == ".csv":
        return _load_payload_from_csv(data)
    if suffix in {".xlsx", ".xls"}:
        return _load_payload_from_excel(data)
    if suffix == ".docx":
        return _load_payload_from_docx(data)
    if suffix == ".pdf":
        return _load_payload_from_pdf(data)
    raise ValueError(f"Unsupported file type: {suffix}")


def load_payload_from_text(text: str) -> Mapping[str, object]:
    stripped = text.strip()
    if not stripped:
        raise ValueError("Uploaded file was empty.")
    try:
        return json.loads(stripped)
    except json.JSONDecodeError as exc:  # pragma: no cover - invalid user input
        raise ValueError("Uploaded document does not contain valid JSON assumptions.") from exc


def extract_text_from_upload(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".csv"}:
        return data.decode("utf-8", errors="replace")
    if suffix == ".docx":
        if Document is None:
            raise ValueError("Word support requires the 'python-docx' package to be installed.")
        document = Document(io.BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
        if not text:
            raise ValueError("Word document did not contain any readable text.")
        return text
    if suffix == ".pdf":
        if PdfReader is None:
            raise ValueError("PDF support requires the 'PyPDF2' package to be installed.")
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
        if not text:
            raise ValueError("PDF file did not contain any readable text.")
        return text
    raise ValueError("Unsupported file type for RAG. Upload TXT, MD, CSV, DOCX, or PDF files.")


def build_rag_chunks(
    documents: List[Mapping[str, str]],
    *,
    size: int,
    overlap: int,
) -> List[Mapping[str, object]]:
    chunks: List[Mapping[str, object]] = []
    for document in documents:
        name = str(document.get("name", "Document"))
        text = str(document.get("text", "") or "").strip()
        for chunk in _chunk_text(text, size=size, overlap=overlap):
            chunks.append({"source": name, "text": chunk, "tokens": _tokenize(chunk)})
    return chunks


def score_chunks(
    query: str,
    chunks: List[Mapping[str, object]],
    *,
    limit: int,
) -> List[Mapping[str, object]]:
    query_tokens = set(_tokenize(query))
    if not query_tokens:
        return []
    scored: List[Mapping[str, object]] = []
    for chunk in chunks:
        tokens = chunk.get("tokens", [])
        overlap = sum(1 for token in tokens if token in query_tokens)
        if overlap == 0:
            continue
        score = overlap / max(len(tokens), 1)
        scored.append({**chunk, "score": score})
    scored.sort(key=lambda item: item.get("score", 0.0), reverse=True)
    return scored[:limit]


def _extract_json_fragment(text: str) -> str:
    if "{" in text and "}" in text:
        start = text.find("{")
        end = text.rfind("}")
        if end > start:
            return text[start : end + 1]
    return text


def _load_payload_from_csv(data: bytes) -> Mapping[str, object]:
    text = data.decode("utf-8-sig")
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        fragment = _extract_json_fragment(text)
        if fragment and fragment != text:
            return load_payload_from_text(fragment)

        reader = csv.reader(io.StringIO(text))
        cells: list[str] = []
        for row in reader:
            cells.extend(cell for cell in row if cell is not None)
        joined = _extract_json_fragment("".join(cells).strip())
        if not joined:
            raise ValueError("CSV file did not contain any usable JSON text.")
        return load_payload_from_text(joined)


def _load_payload_from_excel(data: bytes) -> Mapping[str, object]:
    if load_workbook is None:  # pragma: no cover - optional dependency path
        raise ValueError("Excel support requires the 'openpyxl' package to be installed.")

    workbook = load_workbook(filename=io.BytesIO(data), data_only=True)
    text_parts: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                text_parts.append(str(value))
    combined = _extract_json_fragment("\n".join(text_parts).strip())
    if not combined:
        raise ValueError("Excel file did not contain any readable text.")
    return load_payload_from_text(combined)


def _load_payload_from_docx(data: bytes) -> Mapping[str, object]:
    if Document is None:  # pragma: no cover - optional dependency path
        raise ValueError("Word support requires the 'python-docx' package to be installed.")

    document = Document(io.BytesIO(data))
    text = _extract_json_fragment(
        "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
    )
    if not text:
        raise ValueError("Word document did not contain any readable text.")
    return load_payload_from_text(text)


def _load_payload_from_pdf(data: bytes) -> Mapping[str, object]:
    if PdfReader is None:  # pragma: no cover - optional dependency path
        raise ValueError("PDF support requires the 'PyPDF2' package to be installed.")

    reader = PdfReader(io.BytesIO(data))
    text_parts: list[str] = []
    for page in reader.pages:
        extracted = page.extract_text() or ""
        text_parts.append(extracted)
    combined = _extract_json_fragment("\n".join(text_parts).strip())
    if not combined:
        raise ValueError("PDF file did not contain any readable text.")
    return load_payload_from_text(combined)


def _tokenize(text: str) -> List[str]:
    return re.findall(r"[A-Za-z0-9']+", text.lower())


def _chunk_text(text: str, *, size: int, overlap: int) -> List[str]:
    words = text.split()
    if not words:
        return []
    chunks: List[str] = []
    step = max(1, size - overlap)
    for start in range(0, len(words), step):
        chunk = " ".join(words[start:start + size]).strip()
        if chunk:
            chunks.append(chunk)
    return chunks
